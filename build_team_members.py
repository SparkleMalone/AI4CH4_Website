#!/usr/bin/env python3
"""
Read 'WG Member Contact 2025 web.xlsx' and write team_members.json for the team table page.
Columns: B=Email, C=Name, D=Position, E=Affiliation (1-based).
Emails are deliberately NOT written: team_members.json is publicly served.
"""
import json
import openpyxl

EXCEL_PATH = "WG Member Contact 2025 web.xlsx"
OUTPUT_JSON = "team_members.json"
# 1-based: A=1, B=2, C=3, D=4, E=5
COL_NAME = 3
COL_POSITION = 4
COL_AFFILIATION = 5


def main():
    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True))
    wb.close()

    members = []
    for row in rows:
        if not row or len(row) < COL_AFFILIATION:
            continue
        name = row[COL_NAME - 1]
        if not name:
            continue
        name = str(name).strip() if name else ""
        position = (row[COL_POSITION - 1] or "")
        position = str(position).strip() if position else ""
        affiliation = (row[COL_AFFILIATION - 1] or "")
        affiliation = str(affiliation).strip() if affiliation else ""
        members.append({
            "name": name,
            "position": position,
            "affiliation": affiliation,
        })

    with open(OUTPUT_JSON, "w", encoding="utf-8") as f:
        json.dump(members, f, indent=2, ensure_ascii=False)
    print("Wrote", OUTPUT_JSON, "with", len(members), "team members.")


if __name__ == "__main__":
    main()
