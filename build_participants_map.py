#!/usr/bin/env python3
"""
Read affiliations from 'WG Member Contact 2025 web.xlsx' and write
participants_map_data.json for the world map.

Every institution in KNOWN is placed at fixed, hand-checked coordinates with a
short place label. Only an affiliation that is NOT in KNOWN is sent to the
Nominatim geocoder. Check where it lands, then add it to KNOWN so the result
stays stable on later runs.
"""
import json
import time
from collections import defaultdict

import openpyxl

EXCEL_PATH = "WG Member Contact 2025 web.xlsx"
AFFILIATION_COL = 5  # 1-based column E
OUTPUT_JSON = "participants_map_data.json"
NOMINATIM_DELAY = 1.1  # seconds between requests (rate limit)

BOULDER_NOAA = (39.9956, -105.2620, "Boulder, CO, USA")
COLLEGE_PARK = (38.9869, -76.9426, "College Park, MD, USA")

# affiliation (exactly as written in the spreadsheet) -> (lat, lon, place label)
KNOWN = {
    "AtmoFacts": (40.1672, -105.1019, "Longmont, CO, USA"),
    "CIRES, CU Boulder and NOAA Global Monitoring Laboratory": BOULDER_NOAA,
    "Colorado State University": (40.5734, -105.0865, "Fort Collins, CO, USA"),
    "Harvard University": (42.3770, -71.1167, "Cambridge, MA, USA"),
    "Institute of Tibetan Plateau Research, Chinese Academy of Sciences": (39.9, 116.4, "Beijing, China"),
    "Institute of Tibetan Plateau Research, Chinese Academy of Sciences, Beijing": (39.9, 116.4, "Beijing, China"),
    "Joint Global Change Research Institute, Pacific Northwest National Lab": COLLEGE_PARK,
    "Kangwon National University": (37.8693, 127.7444, "Chuncheon, South Korea"),
    "Lawrence Berkeley National Laboratory": (37.8762, -122.2467, "Berkeley, CA, USA"),
    "McGill University": (45.5048, -73.5772, "Montréal, QC, Canada"),
    "NASA Goddard": (38.9949, -76.8521, "Greenbelt, MD, USA"),
    "NOAA Global Monitoring Laboratory": BOULDER_NOAA,
    "NOAA Global Monitoring Laboratory & CU Boulder": BOULDER_NOAA,
    "National Center for Atmospheric Research": (39.9780, -105.2750, "Boulder, CO, USA"),
    "Northern Arizona University": (35.1894, -111.6513, "Flagstaff, AZ, USA"),
    "Pacific Northwest National Laboratory": (46.3456, -119.2789, "Richland, WA, USA"),
    "Purdue University": (40.4237, -86.9212, "West Lafayette, IN, USA"),
    "Seoul National University": (37.4599, 126.9519, "Seoul, South Korea"),
    "Spark Climate Solutions": (37.77, -122.42, "San Francisco, CA, USA"),
    "Stanford University": (37.4275, -122.1697, "Stanford, CA, USA"),
    "University at Albany": (42.6860, -73.8236, "Albany, NY, USA"),
    "University of Colorado": (40.0076, -105.2659, "Boulder, CO, USA"),
    "University of Houston": (29.7199, -95.3422, "Houston, TX, USA"),
    "University of Illinois Chicago": (41.8708, -87.6505, "Chicago, IL, USA"),
    "University of Maryland, College Park": COLLEGE_PARK,
    "UMD ESSIC & NASA GSFC": COLLEGE_PARK,
    "Earth System Science Interdisciplinary Center, University of Maryland College Park": COLLEGE_PARK,
    "University of Minnesota": (44.9740, -93.2277, "Minneapolis, MN, USA"),
    "University of Pittsburgh": (40.4444, -79.9608, "Pittsburgh, PA, USA"),
    "University of Texas at Austin": (30.2849, -97.7341, "Austin, TX, USA"),
    "University of Wisconsin–Madison": (43.0766, -89.4125, "Madison, WI, USA"),
    "Woodwell Climate Research Center": (41.5490, -70.6440, "Falmouth, MA, USA"),
    "Yale University": (41.3163, -72.9223, "New Haven, CT, USA"),
}


def geocode(aff):
    """Look up an affiliation not in KNOWN. Returns (lat, lon, place) or None."""
    from geopy import Nominatim
    from geopy.exc import GeocoderTimedOut, GeocoderServiceError

    geocoder = Nominatim(user_agent="AI4Methane-participants-map")
    time.sleep(NOMINATIM_DELAY)
    try:
        loc = geocoder.geocode(aff, timeout=10)
    except (GeocoderTimedOut, GeocoderServiceError):
        print("Skip (geocode error):", aff)
        return None
    if not loc:
        print("Skip (no result):", aff)
        return None
    print("Geocoded (please check, then add to KNOWN):", aff, "->", loc.address)
    return loc.latitude, loc.longitude, loc.address


def main():
    # 1) Read affiliations from Excel
    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    ws = wb.active
    affiliations = []
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        if row and len(row) >= AFFILIATION_COL:
            aff = row[AFFILIATION_COL - 1]
            if aff and isinstance(aff, str) and aff.strip():
                affiliations.append(aff.strip())
    wb.close()

    aff_count = defaultdict(int)
    for a in affiliations:
        aff_count[a] += 1

    # 2) Place each unique affiliation; merge pins that fall within ~1 km
    loc_to_data = {}
    for aff, count in sorted(aff_count.items()):
        hit = KNOWN.get(aff) or geocode(aff)
        if not hit:
            continue
        lat, lon, place = hit
        key = (round(lat, 2), round(lon, 2))
        d = loc_to_data.setdefault(
            key, {"lat": lat, "lon": lon, "place": place, "institutions": [], "count": 0}
        )
        d["institutions"].append({"name": aff, "count": count})
        d["count"] += count

    # 3) One entry per location, largest first
    out = sorted(loc_to_data.values(), key=lambda d: (-d["count"], d["place"]))
    with open(OUTPUT_JSON, "w", encoding="utf-8") as f:
        json.dump(out, f, indent=2, ensure_ascii=False)
    print("Wrote", OUTPUT_JSON, "with", len(out), "locations and",
          sum(d["count"] for d in out), "participants.")


if __name__ == "__main__":
    main()
